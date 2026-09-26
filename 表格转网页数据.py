# -*- coding: utf-8 -*-
"""数据源：体测数据.xlsx  →  输出：data.js（网页自动加载）
每一步：在 xlsx 里加行 -> 运行本脚本 -> 刷新网页即更新。
运行： python3 表格转网页数据.py
"""
import openpyxl, json

SRC = "体测数据.xlsx"
OUT = "data.js"
PEOPLE = ("xym", "gcc")          # 与 xlsx 里的工作表同名，按需增删

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0

def person(name):
    ws = wb[name]
    out = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value in (None, ""):   # 跳过空行
            continue
        g = lambda c: num(ws.cell(r, c).value)
        out.append({
            "date": ws.cell(r, 1).value,
            "weight": g(2), "muscle": g(3), "fatMass": g(4), "fatPct": g(5), "water": g(6),
            "protein": g(7), "mineral": g(8), "leanMass": g(9), "bmi": g(10),
            "visceral": g(11), "bmr": g(12),
            "segMuscle":   {"左臂": g(13), "右臂": g(14), "躯干": g(15), "左腿": g(16), "右腿": g(17)},
            "segFat":      {"左臂": g(18), "右臂": g(19), "躯干": g(20), "左腿": g(21), "右腿": g(22)},
            "segMusclePct":{"左臂": g(23), "右臂": g(24), "躯干": g(25), "左腿": g(26), "右腿": g(27)},
            "segFatPct":   {"左臂": g(28), "右臂": g(29), "躯干": g(30), "左腿": g(31), "右腿": g(32)},
            "idealWeight": g(33), "calorie": g(34),
            "goalWeight": g(35), "goalFat": g(36), "goalMuscle": g(37),
        })
    return out

wb = openpyxl.load_workbook(SRC)
data = { name: {"title": name, "records": person(name)} for name in PEOPLE if name in wb.sheetnames }

js  = "/* 本文件由「表格转网页数据.py」从「体测数据.xlsx」自动生成 —— 请勿手改，改数据请改 xlsx 后重新运行脚本。 */\n"
js += "const DATA = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n"
open(OUT, "w", encoding="utf-8").write(js)
print("已生成", OUT, "|", {k: len(v["records"]) for k, v in data.items()})
